"""Convert docs/SETUP_GUIDE.md to a styled Word document (docs/Vyuh_Setup_Guide.docx).

Needs pandoc (https://pandoc.org) and python-docx. Run from the project root:
    python docs/build_docx.py
"""
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC, OUT = ROOT / "SETUP_GUIDE.md", ROOT / "Vyuh_Setup_Guide.docx"


def borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "C9D1DE")
        b.append(el)
    pr.append(b)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def main():
    subprocess.run(["pandoc", str(SRC), "-o", str(OUT)], check=True)
    doc = Document(str(OUT))
    st = doc.styles
    st["Normal"].font.name = "Calibri"; st["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12), ("Title", 24)):
        if name in st:
            st[name].font.name = "Calibri"; st[name].font.size = Pt(size)
            st[name].font.color.rgb = RGBColor(0x24, 0x57, 0xD6)
    for name in ("Source Code", "Verbatim Char"):
        if name in st:
            st[name].font.name = "Consolas"; st[name].font.size = Pt(9)
    for t in doc.tables:
        borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for cell in t.rows[0].cells:
            shade(cell, "E7EEFC")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)
    for p in doc.paragraphs:
        if p.style.name == "Source Code":
            p.paragraph_format.left_indent = Pt(12)
    doc.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    sys.exit(main())
